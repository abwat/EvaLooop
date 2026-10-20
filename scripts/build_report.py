"""Build the three-page report from the accepted experiment results."""

from pathlib import Path
import argparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def build(output):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.65)
    section.footer_distance = Inches(0.25)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.0
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for name, size in [('Title', 17), ('Heading 1', 12)]:
        style = doc.styles[name]
        style.font.name, style.font.size = 'Calibri', Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(4)
    doc.core_properties.title = 'Spanish Summaries in EvaLoop'
    doc.core_properties.author = 'Aniruddh S Bhagwat'

    def p(text, style=None):
        return doc.add_paragraph(text, style)

    def h(text):
        p(text, 'Heading 1')

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        for cell, value in zip(t.rows[0].cells, headers):
            cell.text = value
        repeat = OxmlElement('w:tblHeader')
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells = t.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)
        for i, row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for j, cell in enumerate(row.cells):
                cell.width = Inches(widths[j])
                cell.vertical_alignment = 1
                tcpr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for edge in ['top', 'left', 'bottom', 'right']:
                    b = OxmlElement('w:' + edge)
                    for key, val in [('val','single'),('sz','4'),('color','D9D9D9')]:
                        b.set(qn('w:' + key), val)
                    borders.append(b)
                tcpr.append(borders)
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'E5E7EB' if i == 0 else ('F5F6F7' if i % 2 == 0 else 'FFFFFF'))
                tcpr.append(shading)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(4)
                    paragraph.paragraph_format.space_after = Pt(4)
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5)
                        run.bold = i == 0
        p('').paragraph_format.space_after = Pt(0)

    p('Spanish Summaries in EvaLoop', 'Title')
    sub = p('Aniruddh S Bhagwat | CSC 591/791 | Assignment 4 | 20 October 2026')
    sub.paragraph_format.space_after = Pt(8)
    h('Research question and design')
    p('Does requiring Spanish intermediate summaries reduce robustness after controlling for the summary prompt? '
      'The hypothesis is a negative Spanish-minus-English effect on quadratic Average Sustainable Loops (ASL). '
      'The estimated pooled effect is −0.074 ASL (95% CI [−0.432, 0.294]), with different directions across models. '
      'The experiment therefore does not establish a consistent Spanish penalty.')
    p('This Option B extension compares three conditions on identical tasks: the supplied standard-loop records, '
      'a new English control requesting a detailed summary, and the same control requesting Spanish. The two new '
      'arms differ only in the language substitution. Only the model’s own summary perturbs the next task. '
      'Historical comparisons also change prompt wording and possibly backend settings, so they are descriptive.')
    h('The evaluation pipeline and extension')
    p('Each task starts with its original MBPP+ description and assertion. The package generation template asks '
      '“Generate Python code for the following task” and “Return only the code, without explanations or comments.” '
      'The runner extracts the first fenced code block, or the whole response if no block exists, then calls '
      'EvalPlusCodeTester. Both Base and Plus tests must pass. A failure ends the task immediately; a passing '
      'program is summarized by the same model. The loop ends after the first failure or ten passing cycles.')
    p('The experimental summary prompt begins “Summarize what the following code does in a detailed {language} '
      'description. Preserve every behavioral requirement, edge case, input/output constraint, and function-signature '
      'requirement.” It ends “Return only the summary in {language}, without explanations, comments, or code.” '
      'The code is enclosed in a fenced block; {language} is English or Spanish. The next description is exactly '
      'a triple-quoted block containing the stripped summary followed by the final assertion from the original task. '
      'No previous prose or code is carried forward. The assertion anchors the callable name and example argument '
      'structure, reducing signature drift without guaranteeing semantic correctness. Its English identifiers are '
      'a deliberate language leak.')
    p('New files in evaloop/extensions/ implement the prompts, extraction, wrapper, and metrics. New scripts handle '
      'cloud inference, atomic task checkpoints, analysis, and qualitative codes. The existing generation and test '
      'modules remain unchanged; README changes document the extension. The runner reuses the upstream tester '
      'and MBPP loader. Tests check the copied prompt and wrapper bytes, reconstruct a supplied leaderboard row, '
      'and cover extraction, arm equivalence, first-failure stopping, and ten-cycle termination. The cloud notebook '
      'records 84 passing tests.')
    h('Experimental setup')
    p('The first 100 IDs of a seed-20261017 permutation of all 378 supplied tasks form the sample. The first 25 '
      'are the gate and smoke subset. Models are Qwen2.5-Coder-7B-Instruct, CodeQwen1.5-7B-Chat, and '
      'OpenCoder-8B-Instruct, with immutable revisions in the configuration. This final convenience set fits one '
      '24 GB NVIDIA L4 in BF16. All six new runs use batch size 1, repository chat templates, no system prompt, '
      'greedy decoding (do_sample=False), 1,024 new tokens per generation, and at most ten cycles.')
    p('The archived protocol already specifies Lingua and the final model set; earlier discussion considered '
      'Gemma and cycle-count tests. No dated change log verifies those planning transitions. Full settings, '
      'provenance, audit selection, cost qualifications, and rebuild commands are documented in the README '
      'and POST_RUN_AUDIT.md. The added audits and sensitivity checks are exploratory.')

    doc.add_page_break()
    h('Checks and results')
    p('The Qwen2.5 gate repeated 25 identical cycle-1 prompts three times and recorded no pass/fail disagreement. '
      'Across the 100-task English and Spanish arms, cycle-1 disagreement is 0/100 for each model. Equal success '
      'rates are a noise check because the initial prompts are identical. English-control agreement with supplied '
      'cycle-1 outcomes is 100/100 for CodeQwen, 95/100 for OpenCoder, and 92/100 for Qwen2.5; the historical '
      'baseline is therefore not an exact execution reproduction.')
    p('OpenCoder required its pinned repository code and generation-config EOS token. Its affected smoke outputs '
      'were discarded and rerun after that correction. Other smoke outputs were reused only when settings were '
      'unchanged. Token-cap counts (English, Spanish) are CodeQwen 0/1,294 and 21/1,364; OpenCoder 1/1,262 '
      'and 2/1,229; Qwen2.5 0/1,334 and 0/1,268. Excluding pairs with any cap leaves 97 CodeQwen tasks '
      'and an effect of +0.411 versus +0.399 originally. This post hoc exclusion preserves the sign but cannot '
      'identify cap-induced bias. Lingua’s binary English/Spanish detector marks every Spanish summary compliant. '
      'An AI-assisted reading of 30 summaries (ten distinct task IDs per model, including short and '
      'identifier-heavy prose) finds 29 Spanish-only prose samples and one with the English word “effectively.” '
      'All 30 are predominantly Spanish; this is not independent human validation.')
    table(['Model', 'Success\nStd / Eng / Spa', 'ASL\nStandard', 'ASL\nEnglish', 'ASL\nSpanish'], [
        ['CodeQwen1.5 7B', '.70 / .70 / .70', '6.313', '6.208', '6.607'],
        ['OpenCoder 8B', '.67 / .68 / .68', '6.304', '6.019', '5.789'],
        ['Qwen2.5 Coder 7B', '.72 / .70 / .70', '6.333', '6.417', '6.026'],
    ], [1.6, 1.8, 1.2, 1.2, 1.2])
    p('Standard values are recomputed from supplied records on the same 100 IDs. If s is the number of passing '
      'cycles (0–10), success rate is mean(s ≥ 1) and quadratic ASL is mean(s²/10). Neither is a judge score.')
    table(['Model', 'Spanish − English\nASL', '95% bootstrap CI', 'Wilcoxon\np', 'Shorter / equal /\nlonger loops'], [
        ['CodeQwen1.5', '+0.399', '[0.000, 0.891]', '.098', '1 / 94 / 5'],
        ['OpenCoder', '−0.230', '[−0.837, 0.380]', '.263', '8 / 88 / 4'],
        ['Qwen2.5 Coder', '−0.391', '[−1.095, 0.294]', '.203', '10 / 85 / 5'],
        ['Pooled', '−0.074', '[−0.432, 0.294]', '—', '19 / 267 / 14'],
    ], [1.4, 1.3, 1.8, .8, 1.7])
    p('Each interval uses 10,000 percentile bootstrap resamples of tasks. The paired, two-sided Wilcoxon test '
      'uses per-task s²/10 differences and Pratt zero handling. The secondary pooled analysis averages the three '
      'fixed models within each task before resampling tasks, preserving their shared task dependence. Pooled '
      'counts are 300 model-task pairs, not 300 independent tasks. Cycle-count differences give the same p-values. '
      'All intervals include or touch zero and all p-values exceed .05. The intervals still allow roughly a 7% '
      'pooled loss or a 17% Qwen2.5 loss relative to English ASL; this is not an equivalence result.')
    p('English-control ASL ranks Qwen2.5 > CodeQwen > OpenCoder; Spanish ranks CodeQwen > Qwen2.5 > OpenCoder. '
      'The supplied standard scores are almost tied. With only three models and wide intervals, the change is '
      'descriptive. Most paired tasks are unchanged, and the evidence does not support a stable ranking reversal.')

    doc.add_page_break()
    h('Failure analysis and comparison with the standard loop')
    p('I manually reviewed the 20 Spanish failure samples on 20 October 2026 and retained the draft labels. '
      'The sample was selected with seeded stratification over models and '
      'failure cycles. A post hoc audit adds ten English-control and ten standard-loop failures using the same '
      'rule. All samples exclude cycle-1 failures and ten-cycle survivors. One primary label is assigned from '
      'the original prompt, preceding summary, and failing code, consulting the preceding passing code as needed.')
    table(['Condition', 'n', 'Requirement\ndropped', 'Wrong\nimplementation', 'Other\nprimary'], [
        ['Spanish', '20', '11 (55%)', '9 (45%)', '0'],
        ['English control', '10', '5 (50%)', '5 (50%)', '0'],
        ['Standard supplied', '10', '5 (50%)', '5 (50%)', '0'],
    ], [1.8, .5, 1.6, 1.7, 1.4])
    p('Both mechanisms occur in every condition. For example, standard CodeQwen MBPP/14 changes a triangular '
      'prism to a trapezoidal prism; standard Qwen2.5 MBPP/227 retains “minimum” but mishandles ties. Repeated '
      'fences in standard OpenCoder MBPP/123 are secondary to incorrect divisor code. Unequal model composition, '
      'small samples, and subjective labels prevent interpreting these proportions as population differences.')
    p('Codex (GPT-6) assisted with draft failure labels and the language audit. My review confirms the 20 '
      'Spanish cases; the supplementary baseline and language labels remain unconfirmed by a human. '
      'Prompts and decision rules are in POST_RUN_AUDIT.md. Exact snapshot and decoding settings are unavailable. '
      'Incremental API cost was $0; subscription cost per example is unallocated. Labels do not affect ASL or tests.')
    h('Representative broken and surviving loops')
    p('Broken case, OpenCoder MBPP/14. The task requires triangular-prism volume, (base × height × length)/2. '
      'The first Spanish summary describes a parallelepiped and omits the half-factor. Cycle 2 returns base × '
      'height × length and fails. English control retains the half-factor despite calling the object a rectangular '
      'box, and both English control and the supplied standard loop survive ten cycles.')
    p('Surviving case, MBPP/101. All nine model-condition combinations survive ten cycles on the kth-element '
      'task. Spanish summaries retain the one-based rule, return arr[k−1], and the valid-range assumption. '
      'The assertion anchors the callable signature.')
    h('Limitations and implications')
    p('Cycle-1 agreement does not measure later-cycle inference noise. The assertion can repair descriptions '
      'or encourage overfitting. Tests may miss semantic errors, quadratic ASL emphasizes long survivors, and '
      'language detection can mislabel mixed prose. Historical backend metadata are missing, limiting baseline '
      'attribution. Public MBPP tasks may appear in training data. Shared Qwen lineage could correlate '
      'memorization and errors, so two Qwen models provide limited independent family evidence. Results do not '
      'generalize beyond these short tasks, models, or Spanish. The observed effects have mixed signs.')
    h('Runtime cost and reproducibility')
    p('Accepted generation time is 11.22 hours on a g2-standard-4 L4 VM, initially Spot and later standard. '
      'The included billing screenshot shows US$20.75 for 14–20 October, exceeding the planned $10 budget. '
      'I forgot to turn off the VM after the experiment, so it kept running and added charges. '
      'The image lacks project filters and does not separate experiment costs from idle charges. '
      'Generated code ran through EvalPlus on the cloud VM, separate from local personal files. '
      'Timers exclude loading, idle time, disks, and interrupted work. Six-run summary.csv rebuilds '
      'without course data using python scripts/analyze_assignment4.py --summary-only; only the historical '
      'baseline rebuild requires --data-dir. Supplied traces are referenced rather than copied.')
    refs = p('References  EvaLoop, arXiv:2505.12185; github.com/evalooop/EvaLooop. EvalPlus, '
             'github.com/evalplus/evalplus. Course data README and prompt templates, accessed 17 October 2026.')
    for run in refs.runs:
        run.font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('Page ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.save(output)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, default=Path('report.docx'))
    build(parser.parse_args().output)
